"""Rapport de corroboration exportable : Excel multi-onglets (mis en forme) et CSV."""
from __future__ import annotations

import io
import json
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .engine import Result
from .models import ANOMALIE, CONFORME, JUSTIFIE

COLUMNS = [
    ("priorite", "Priorité"), ("verdict", "Verdict"), ("a_valider", "À valider"), ("matricule", "Matricule"),
    ("employe", "Employé"), ("type_affectation", "Type aff."), ("code_poste", "Poste"), ("code_emploi", "Emploi"),
    ("description", "Description"), ("champ_a", "Champ Système A"), ("champ_b", "Champ Système B"),
    ("valeur_a", "Valeur A (source)"), ("valeur_attendue", "Valeur attendue (règle)"), ("valeur_b", "Valeur B (cible)"),
    ("niveau", "Niveau de décision"), ("decide_par", "Décidé par"), ("regle_id", "Règle"), ("confiance", "Confiance"),
    ("justification", "Justification"), ("diagnostic", "Diagnostic / cause probable"), ("regle_texte", "Texte de la règle"),
    ("preuves", "Preuves (JSON)"), ("ia", "Détail IA (JSON)"), ("signature", "Motif"), ("id", "Identifiant"),
]
FILL = {ANOMALIE: "F8D7DA", JUSTIFIE: "FFF3CD", CONFORME: "D1E7DD"}


def to_frame(findings) -> pd.DataFrame:
    rows = []
    for f in findings:
        d = f.to_dict()
        d["preuves"] = json.dumps(d["preuves"], ensure_ascii=False, default=str)
        d["ia"] = json.dumps(d["ia"], ensure_ascii=False, default=str) if d["ia"] else ""
        d["a_valider"] = "Oui" if d["a_valider"] else ""
        d["confiance"] = round(d["confiance"], 2)
        rows.append({label: d[key] for key, label in COLUMNS})
    return pd.DataFrame(rows, columns=[c[1] for c in COLUMNS])


def _style(ws, verdict_col: int | None = None, widths: dict | None = None):
    ws.freeze_panes = "A2"
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, col in enumerate(ws.columns, 1):
        header = str(col[0].value or "")
        w = (widths or {}).get(header) or min(max(len(header), *(len(str(c.value or "")) for c in col[1:50])) + 2, 60)
        ws.column_dimensions[get_column_letter(i)].width = max(w, 10)
    if verdict_col:
        for row in ws.iter_rows(min_row=2):
            v = row[verdict_col - 1].value
            if v in FILL:
                row[verdict_col - 1].fill = PatternFill("solid", fgColor=FILL[v])
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def build_excel(res: Result) -> bytes:
    df = to_frame(res.findings)
    m = res.meta
    c = res.counts()
    synth = m.get("synthese", {})
    summary = [
        ("Rapport de corroboration CorroborIA", ""),
        ("Date d'exécution", m.get("date_execution")),
        ("Durée (s)", m.get("duree_s")),
        ("", ""),
        ("Constats évalués (enregistrement × champ)", c["total"]),
        ("Anomalies réelles à investiguer", c[ANOMALIE]),
        ("Écarts justifiés automatiquement", c[JUSTIFIE]),
        ("Conformes", c[CONFORME]),
        ("Constats à valider par un expert", c["a_valider"]),
        ("", ""),
        ("Lignes Système A / Système B", f"{m['nb_lignes']['source']} / {m['nb_lignes']['destination']}"),
        ("Appariement", ", ".join(f"{k} : {v}" for k, v in m["appariement"].items())),
        ("Intégrité des fichiers sources (sha256)", "OK — fichiers inchangés" if all(i["intact"] for i in m["integrite"]) else "ÉCART DÉTECTÉ"),
        ("Fournisseur(s) IA utilisés", ", ".join(m["llm"]["utilise"]) or "aucun cas ambigu"),
        ("Modèle ML", f"{m['ml']['modele']} — {m['ml']['exemples_entrainement']} exemples, {m['ml']['corrections_expert']} corrections expert"),
        ("", ""),
        ("Synthèse (IA — " + synth.get("fournisseur", "") + ")", synth.get("resume", "")),
    ]
    for cause in synth.get("causes_racines", []):
        summary.append((f"Cause : {cause.get('cause')}", f"{cause.get('nombre')} cas — {', '.join(cause.get('champs', []))} — {cause.get('recommandation')}"))
    for cal in m.get("calibration_regles", []):
        summary.append((f"Calibration {cal['regle']}", f"{cal['interpretations']} → retenue : {cal['retenue']}. {cal['conclusion']}"))

    per_field = (df.groupby(["Champ Système B", "Verdict"]).size().unstack(fill_value=0)
                 .reindex(columns=[ANOMALIE, JUSTIFIE, CONFORME], fill_value=0).reset_index())
    rules = pd.DataFrame([{"Règle": k, "Description": v} for k, v in m["catalogue_regles"].items()])
    coverage = pd.DataFrame(m["couverture_mapping"])
    integrity = pd.DataFrame(m["integrite"])
    providers = pd.DataFrame(m["llm"]["fournisseurs"])

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        pd.DataFrame(summary, columns=["Indicateur", "Valeur"]).to_excel(xw, sheet_name="Synthèse", index=False)
        sheets = [("Anomalies", df[df["Verdict"] == ANOMALIE]), ("Écarts justifiés", df[df["Verdict"] == JUSTIFIE]),
                  ("Conformes", df[df["Verdict"] == CONFORME]), ("Tous les constats", df)]
        for name, part in sheets:
            part.to_excel(xw, sheet_name=name, index=False)
        per_field.to_excel(xw, sheet_name="Par champ", index=False)
        rules.to_excel(xw, sheet_name="Règles", index=False)
        coverage.to_excel(xw, sheet_name="Couverture mapping", index=False)
        providers.to_excel(xw, sheet_name="Fournisseurs IA", index=False)
        integrity.to_excel(xw, sheet_name="Intégrité", index=False)
        wb = xw.book
        _style(wb["Synthèse"], widths={"Indicateur": 45, "Valeur": 120})
        for row in wb["Synthèse"].iter_rows(min_row=2):
            row[1].alignment = Alignment(wrap_text=True, vertical="top")
        wb["Synthèse"]["A1"].value = "Indicateur"
        verdict_idx = [c[1] for c in COLUMNS].index("Verdict") + 1
        wide = {"Justification": 70, "Diagnostic / cause probable": 60, "Texte de la règle": 50, "Preuves (JSON)": 50,
                "Détail IA (JSON)": 40, "Valeur A (source)": 35}
        for name, _ in sheets:
            _style(wb[name], verdict_idx, wide)
        for name in ("Par champ", "Règles", "Couverture mapping", "Fournisseurs IA", "Intégrité"):
            _style(wb[name])
    return buf.getvalue()


def build_csv(res: Result, only: str | None = None) -> bytes:
    df = to_frame(res.findings)
    if only:
        df = df[df["Verdict"] == only]
    return df.to_csv(index=False, sep=";").encode("utf-8-sig")


def _md(v, n: int = 160) -> str:
    s = "∅" if v in (None, "") else str(v).replace("\n", " ").replace("|", "\\|")
    return s if len(s) <= n else s[: n - 1] + "…"


def build_markdown(res: Result) -> str:
    """Résultats finaux lisibles directement sur GitHub (sans exécuter le code ni ouvrir Excel)."""
    m, c = res.meta, res.counts()
    anomalies = [f for f in res.findings if f.verdict == ANOMALIE]
    justified = [f for f in res.findings if f.verdict == JUSTIFIE]
    synth = m.get("synthese", {})
    intact = all(i["intact"] for i in m["integrite"])
    L = [
        "# Résultats finaux de la corroboration — CorroborIA",
        "",
        f"Généré le {m['date_execution']} à partir des extractions fournies par Loto-Québec (fichiers sources inchangés : "
        f"{'✓ sha256 conformes' if intact else '⚠ écart sha256'}). LLM : {', '.join(m['llm']['utilise']) or 'aucun'}.",
        "Fichiers détaillés : [`rapport_corroboration.xlsx`](rapport_corroboration.xlsx) (10 onglets), "
        "[`rapport_corroboration.csv`](rapport_corroboration.csv) (tous les constats), [`anomalies.csv`](anomalies.csv).",
        "",
        "## Synthèse",
        "",
        "| Verdict | Nombre |",
        "|---|---|",
        f"| Constats évalués (affectations × champs du mapping) | **{c['total']}** |",
        f"| Anomalies | **{c[ANOMALIE]}** (dont {sum(f.a_valider for f in anomalies)} à valider par un expert) |",
        f"| Écarts justifiés automatiquement | **{c[JUSTIFIE]}** |",
        f"| Conformes | **{c[CONFORME]}** |",
        "",
    ]
    if synth.get("resume"):
        L += [f"> {synth['resume']}", ""]
    if synth.get("causes_racines"):
        L += ["### Causes racines (par ordre d'importance)", "", "| Cause | Cas | Recommandation |", "|---|---|---|"]
        L += [f"| {_md(cr.get('cause'), 120)} | {cr.get('nombre')} | {_md(cr.get('recommandation'), 220)} |" for cr in synth["causes_racines"]]
        L.append("")
    L += ["## Anomalies (triées par priorité)", "",
          "| Priorité | Matricule | Aff. | Champ (B) | Attendu (règle) | Reçu (B) | Règle | À valider | Cause probable |",
          "|---|---|---|---|---|---|---|---|---|"]
    L += [f"| {f.priorite} | {f.matricule} | {f.type_affectation or ''} | `{f.champ_b}` | {_md(f.valeur_attendue, 40)} | "
          f"{_md(f.valeur_b, 40)} | {f.regle_id} | {'oui' if f.a_valider else ''} | {_md(f.diagnostic or f.justification, 200)} |"
          for f in anomalies]
    L += ["", "## Écarts justifiés (regroupés)", "", "| Champ (B) | Règle | Cas | Exemple | Justification |", "|---|---|---|---|---|"]
    groups: dict[tuple, list] = {}
    for f in justified:
        groups.setdefault((f.champ_b, f.regle_id), []).append(f)
    for (champ, regle), items in sorted(groups.items(), key=lambda kv: -len(kv[1])):
        ex = items[0]
        L.append(f"| `{champ}` | {regle} | {len(items)} | {ex.matricule} : {_md(ex.valeur_a, 50)} → {_md(ex.valeur_b, 50)} | "
                 f"{_md(ex.justification, 200)} |")
    L += ["", "## Calibration de règle", ""]
    for cal in m.get("calibration_regles", []):
        L.append(f"- **{cal['regle']}** — {cal['texte_mapping']} : " + ", ".join(f"{k} **{v}**" for k, v in cal["interpretations"].items())
                 + f". {cal['conclusion']}")
    L += ["", "## Précisions obtenues de Loto-Québec et prises en compte", "",
          "- **Courriel** : le « code » est le matricule ; le préfixe d'environnement (`dev-08-v2_`) est optionnel côté destination "
          "et accepté. Les adresses utilisant un autre identifiant sont une erreur d'anonymisation du jeu de test, conservée "
          "dans la détection à leur demande (faible priorité).",
          "- **Libellé de rôle** : un préfixe différent du code emploi est une erreur, signalée comme écart (faible priorité).",
          "- **Date d'effet de l'affectation** : la destination applique la règle transformée ; les écarts correspondants sont justifiés.",
          "", "## Méthode (rappel)", "",
          "Niveau 1 : comparaison normalisée · Niveau 2 : règles métier du `Mapping.xlsx` (jointures Motif et Détail du poste) · "
          "Niveau 3 : IA pour les cas ambigus (détecteurs de motifs, scikit-learn, Gemini) · Expert : corrections apprises. "
          "Détails : [README](../README.md) · démo en ligne : https://louisbarbonet.github.io/CorroborAI/", ""]
    return "\n".join(L)


def export(res: Result, out_dir: Path) -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = {"excel": out_dir / "rapport_corroboration.xlsx", "csv": out_dir / "rapport_corroboration.csv",
             "csv_anomalies": out_dir / "anomalies.csv", "resultats_md": out_dir / "RESULTATS.md"}
    paths["excel"].write_bytes(build_excel(res))
    paths["csv"].write_bytes(build_csv(res))
    paths["csv_anomalies"].write_bytes(build_csv(res, ANOMALIE))
    paths["resultats_md"].write_text(build_markdown(res), encoding="utf-8")
    return paths
